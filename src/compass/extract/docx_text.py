"""Extract text from a resume DOCX file."""

from __future__ import annotations

from pathlib import Path

import docx


def extract_text(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)

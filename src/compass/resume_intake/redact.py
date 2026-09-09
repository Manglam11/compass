"""Redact suspected PII from resumes into a separate output directory.

PDF redaction uses PyMuPDF's true redaction (add_redact_annot +
apply_redactions), which strips the underlying content — not a drawn box,
which would leave the original text extractable underneath.

DOCX redaction rewrites the runs that make up each paragraph (and table
cell, and header/footer paragraph and cell) in place. Detection runs on the
concatenated paragraph text — which includes the text of any hyperlinks —
so a match that Word split across run boundaries (common when formatting
changes mid-word) or that sits entirely inside a hyperlink's display-text
run (which is not a direct child of the paragraph and so invisible to
`paragraph.runs`) is still found; matches are then mapped back onto the
run(s) that make them up. A hyperlink's *target* (the relationship, e.g. a
mailto: link or a personal-site URL) is checked and redacted independently
of its display text, since a redacted display label can still point at a
live leak underneath. Names are redacted too, but only when a per-resume
`name_patterns` mapping (from `compass.resume_intake.names`) has an entry
for that file's resume_id (its path stem) — there is no automatic name
detection, and a resume with no entry gets no name redaction.

PDF redaction batches annotations per page: all redaction rects for a page
are collected first, `add_redact_annot` is called for each, and
`apply_redactions()` is called exactly once per page afterward. PyMuPDF's
annotation-id assignment scans existing annotations on every
`add_redact_annot` call, so redundantly searching for and re-adding rects
for a value already handled on that page is quadratic — each unique
matched value is therefore searched for at most once per page. A page
that still yields an implausible number of redaction rects (default
guard: 200) aborts instead of proceeding, since that volume means a
pattern is over-matching and redacting it would silently destroy the
document.
"""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import docx
import pymupdf
from docx.text.hyperlink import Hyperlink

from compass.resume_intake.pii import PLACEHOLDER, find_pii
from compass.resume_intake.text import SUPPORTED_EXTENSIONS

MAX_REDACT_RECTS_PER_PAGE = 200


class SameDirectoryError(Exception):
    pass


class RedactionOvermatchError(Exception):
    pass


def _scan_pdf_counts(path: Path, name_pattern: re.Pattern[str] | None) -> Counter[str]:
    counts: Counter[str] = Counter()
    doc = pymupdf.open(path)
    try:
        for page in doc:
            for match in find_pii(page.get_text(), name_pattern):
                counts[match.kind] += 1
    finally:
        doc.close()
    return counts


def _redact_pdf(
    input_path: Path, output_path: Path, name_pattern: re.Pattern[str] | None
) -> Counter[str]:
    resume_id = input_path.stem
    counts: Counter[str] = Counter()
    doc = pymupdf.open(input_path)
    try:
        for page in doc:
            matches = find_pii(page.get_text(), name_pattern)
            if not matches:
                continue

            # Search for each distinct matched value at most once per page —
            # calling search_for once per match instance (rather than once
            # per unique value) was the quadratic blow-up: a common token
            # matched repeatedly on a page would be re-searched, and its
            # already-found rects re-added, once for every one of its own
            # occurrences.
            rects_by_value: dict[str, list] = {}
            placeholder_by_value: dict[str, str] = {}
            total_rects = 0
            for match in matches:
                if match.value in rects_by_value:
                    continue
                rects = page.search_for(match.value)
                if not rects:
                    continue
                rects_by_value[match.value] = rects
                placeholder_by_value[match.value] = PLACEHOLDER[match.kind]
                total_rects += len(rects)

            if total_rects > MAX_REDACT_RECTS_PER_PAGE:
                raise RedactionOvermatchError(
                    f"{resume_id} page {page.number + 1}: {total_rects} redaction rects "
                    f"exceeds guard of {MAX_REDACT_RECTS_PER_PAGE} — a pattern is over-matching"
                )

            for match in matches:
                if match.value in rects_by_value:
                    counts[match.kind] += 1

            for value, rects in rects_by_value.items():
                placeholder = placeholder_by_value[value]
                for rect in rects:
                    page.add_redact_annot(rect, text=placeholder, fill=(0, 0, 0))

            page.apply_redactions()
        doc.save(output_path)
    finally:
        doc.close()
    return counts


def _table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested_table in cell.tables:
                yield from _table_paragraphs(nested_table)


def _story_paragraphs(story):
    yield from story.paragraphs
    for table in story.tables:
        yield from _table_paragraphs(table)


def _document_paragraphs(document: docx.document.Document):
    """Every paragraph in the body, tables (incl. nested), and headers/footers.

    Header/footer parts are frequently shared across sections (a section
    "linked to previous" reuses the prior section's header/footer part), so
    parts already visited are skipped to avoid scanning or redacting the
    same paragraph twice.
    """
    yield from _story_paragraphs(document)

    seen_parts = set()
    for section in document.sections:
        for story in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            part = story.part
            if part in seen_parts:
                continue
            seen_parts.add(part)
            yield from _story_paragraphs(story)


def _run_atoms(paragraph):
    """Runs in document order, including runs nested inside hyperlinks.

    A hyperlink's display-text runs are children of `<w:hyperlink>`, not
    direct children of `<w:p>`, so `paragraph.runs` alone would skip them.
    """
    atoms = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            atoms.extend(item.runs)
        else:
            atoms.append(item)
    return atoms


def _redact_hyperlink_target(
    hyperlink: Hyperlink, counts: Counter[str], name_pattern: re.Pattern[str] | None
) -> None:
    address = hyperlink.address
    if not address:
        return
    matches = find_pii(address, name_pattern)
    if not matches:
        return
    new_address = address
    for match in sorted(matches, key=lambda m: -m.start):
        new_address = (
            new_address[: match.start] + PLACEHOLDER[match.kind] + new_address[match.end :]
        )
        counts[match.kind] += 1
    rId = hyperlink._element.rId
    if rId:
        hyperlink.part.rels[rId]._target = new_address


def _scan_paragraph(paragraph, counts: Counter[str], name_pattern: re.Pattern[str] | None) -> None:
    for match in find_pii(paragraph.text, name_pattern):
        counts[match.kind] += 1
    for hyperlink in paragraph.hyperlinks:
        for match in find_pii(hyperlink.address, name_pattern):
            counts[match.kind] += 1


def _scan_docx_counts(path: Path, name_pattern: re.Pattern[str] | None) -> Counter[str]:
    counts: Counter[str] = Counter()
    document = docx.Document(str(path))
    for paragraph in _document_paragraphs(document):
        _scan_paragraph(paragraph, counts, name_pattern)
    return counts


def _redact_paragraph(
    paragraph, counts: Counter[str], name_pattern: re.Pattern[str] | None
) -> None:
    runs = _run_atoms(paragraph)
    if runs:
        texts = [run.text for run in runs]
        full_text = "".join(texts)
        matches = find_pii(full_text, name_pattern)

        if matches:
            offsets: list[tuple[int, int]] = []
            pos = 0
            for t in texts:
                offsets.append((pos, pos + len(t)))
                pos += len(t)

            # Rightmost matches first: within any single run, replacing a
            # later span never shifts the character positions of an
            # earlier one, so already-computed offsets for not-yet-applied
            # matches stay valid.
            for match in sorted(matches, key=lambda m: -m.start):
                placeholder_inserted = False
                for i, (run_start, run_end) in enumerate(offsets):
                    if run_end <= match.start or run_start >= match.end:
                        continue
                    local_start = max(match.start, run_start) - run_start
                    local_end = min(match.end, run_end) - run_start
                    text = runs[i].text
                    replacement = PLACEHOLDER[match.kind] if not placeholder_inserted else ""
                    runs[i].text = text[:local_start] + replacement + text[local_end:]
                    placeholder_inserted = True
                counts[match.kind] += 1

    for hyperlink in paragraph.hyperlinks:
        _redact_hyperlink_target(hyperlink, counts, name_pattern)


def _redact_docx(
    input_path: Path, output_path: Path, name_pattern: re.Pattern[str] | None
) -> Counter[str]:
    counts: Counter[str] = Counter()
    document = docx.Document(str(input_path))
    for paragraph in _document_paragraphs(document):
        _redact_paragraph(paragraph, counts, name_pattern)
    document.save(output_path)
    return counts


def scan_directory(
    directory: Path, name_patterns: dict[str, re.Pattern[str]] | None = None
) -> dict[str, Counter[str]]:
    """Report suspected-PII counts per file without modifying anything.

    `name_patterns` maps resume_id (a file's stem, e.g. "prac_001") to a
    pattern scoped to that resume only — a resume_id absent from the
    mapping gets no name matching. Uses the same detection paths
    (paragraphs, tables, headers/footers, hyperlink targets) that
    redaction itself uses, so this is a faithful preview of what
    redaction would find — and, run against an already redacted
    directory, a faithful check of what it missed.
    """
    paths = sorted(p for p in directory.iterdir() if p.suffix.lower() in SUPPORTED_EXTENSIONS)
    results: dict[str, Counter[str]] = {}
    for path in paths:
        is_pdf = path.suffix.lower() == ".pdf"
        name_pattern = name_patterns.get(path.stem) if name_patterns else None
        results[path.name] = (
            _scan_pdf_counts(path, name_pattern)
            if is_pdf
            else _scan_docx_counts(path, name_pattern)
        )
    return results


def redact_directory(
    input_dir: Path,
    output_dir: Path,
    dry_run: bool,
    name_patterns: dict[str, re.Pattern[str]] | None = None,
) -> dict[str, Counter[str]]:
    if input_dir.resolve() == output_dir.resolve():
        raise SameDirectoryError("--output-dir must be different from --input-dir")

    if dry_run:
        return scan_directory(input_dir, name_patterns)

    paths = sorted(p for p in input_dir.iterdir() if p.suffix.lower() in SUPPORTED_EXTENSIONS)
    results: dict[str, Counter[str]] = {}
    output_dir.mkdir(parents=True, exist_ok=True)

    for path in paths:
        is_pdf = path.suffix.lower() == ".pdf"
        output_path = output_dir / path.name
        name_pattern = name_patterns.get(path.stem) if name_patterns else None
        results[path.name] = (
            _redact_pdf(path, output_path, name_pattern)
            if is_pdf
            else _redact_docx(path, output_path, name_pattern)
        )

    return results

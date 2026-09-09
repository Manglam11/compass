"""Tests for compass.resume_intake.redact.

No real resume is used anywhere here — every fixture is built in tmp_path.
"""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from compass.extract.pdf_text import extract_text as extract_pdf_text
from compass.resume_intake.names import compile_name_patterns
from compass.resume_intake.redact import (
    MAX_REDACT_RECTS_PER_PAGE,
    RedactionOvermatchError,
    SameDirectoryError,
    redact_directory,
    scan_directory,
)


def _make_pdf(path: Path, lines: list[str]) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + i * 20), line)
    doc.save(path)
    doc.close()


def _make_docx(path: Path, paragraphs: list[str]) -> None:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)


def _add_hyperlink(paragraph, target: str, display_text: str):
    """Append a `<w:hyperlink>` run to `paragraph`, python-docx has no public API for it."""
    part = paragraph.part
    r_id = part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = display_text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def test_refuses_when_input_and_output_dirs_match(tmp_path: Path):
    with pytest.raises(SameDirectoryError):
        redact_directory(tmp_path, tmp_path, dry_run=True)


def test_dry_run_writes_nothing(tmp_path: Path):
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    _make_docx(input_dir / "resume.docx", ["Email me at jane.doe@example.com please."])

    results = redact_directory(input_dir, output_dir, dry_run=True)

    assert results["resume.docx"]["email"] == 1
    assert not output_dir.exists()
    assert list(input_dir.iterdir()) == [input_dir / "resume.docx"]


def test_docx_email_is_replaced(tmp_path: Path):
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    _make_docx(
        input_dir / "resume.docx",
        ["Experienced engineer.", "Email me at jane.doe@example.com please."],
    )

    results = redact_directory(input_dir, output_dir, dry_run=False)

    assert results["resume.docx"]["email"] == 1
    redacted_path = output_dir / "resume.docx"
    assert redacted_path.exists()

    redacted_document = docx.Document(str(redacted_path))
    full_text = "\n".join(p.text for p in redacted_document.paragraphs)
    assert "jane.doe@example.com" not in full_text
    assert "[EMAIL]" in full_text

    # the input file itself must be untouched
    original_document = docx.Document(str(input_dir / "resume.docx"))
    original_text = "\n".join(p.text for p in original_document.paragraphs)
    assert "jane.doe@example.com" in original_text


def test_pdf_redaction_removes_email_from_text_layer(tmp_path: Path):
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    pdf_path = input_dir / "resume.pdf"
    _make_pdf(
        pdf_path,
        [
            "Experienced Python and SQL developer.",
            "Contact: jane.doe@example.com for details.",
            "Skilled in Docker, AWS, and PostgreSQL.",
        ],
    )

    results = redact_directory(input_dir, output_dir, dry_run=False)

    assert results["resume.pdf"]["email"] == 1
    redacted_path = output_dir / "resume.pdf"
    assert redacted_path.exists()

    redacted_text = extract_pdf_text(redacted_path)
    assert "jane.doe@example.com" not in redacted_text
    assert "[EMAIL]" in redacted_text

    # the input file itself must be untouched
    original_text = extract_pdf_text(pdf_path)
    assert "jane.doe@example.com" in original_text


def test_docx_hyperlink_display_text_and_target_are_both_redacted(tmp_path: Path):
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    document = docx.Document()
    paragraph = document.add_paragraph("Contact: ")
    _add_hyperlink(paragraph, "mailto:jane.doe@example.com", "jane.doe@example.com")
    document.save(input_dir / "resume.docx")

    results = redact_directory(input_dir, output_dir, dry_run=False)
    assert sum(results["resume.docx"].values()) == 2  # display text (email) + target (mailto)

    redacted_document = docx.Document(str(output_dir / "resume.docx"))
    paragraph = redacted_document.paragraphs[0]
    assert "jane.doe@example.com" not in paragraph.text
    assert "[EMAIL]" in paragraph.text
    assert len(paragraph.hyperlinks) == 1
    assert "jane.doe@example.com" not in paragraph.hyperlinks[0].address
    assert paragraph.hyperlinks[0].address == "[EMAIL]"

    # a redacted follow-up scan of the output must find nothing left
    assert sum(scan_directory(output_dir)["resume.docx"].values()) == 0


def test_docx_email_split_across_three_runs_is_redacted(tmp_path: Path):
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Email: jane.")
    paragraph.add_run("doe@exam")
    paragraph.add_run("ple.com done")
    document.save(input_dir / "resume.docx")

    results = redact_directory(input_dir, output_dir, dry_run=False)
    assert results["resume.docx"]["email"] == 1

    redacted_document = docx.Document(str(output_dir / "resume.docx"))
    full_text = redacted_document.paragraphs[0].text
    assert "jane.doe@example.com" not in full_text
    assert "[EMAIL]" in full_text
    assert "done" in full_text


def test_docx_pii_in_table_cell_is_redacted(tmp_path: Path):
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    document = docx.Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Phone: 9876543210"
    document.save(input_dir / "resume.docx")

    results = redact_directory(input_dir, output_dir, dry_run=False)
    assert results["resume.docx"]["phone"] == 1

    redacted_document = docx.Document(str(output_dir / "resume.docx"))
    cell_text = redacted_document.tables[0].cell(0, 0).text
    assert "9876543210" not in cell_text
    assert "[PHONE]" in cell_text


def test_docx_pii_in_header_is_redacted(tmp_path: Path):
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    document = docx.Document()
    document.add_paragraph("Experienced engineer.")
    header = document.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "Jane Doe | jane.doe@example.com"
    document.save(input_dir / "resume.docx")

    results = redact_directory(input_dir, output_dir, dry_run=False)
    assert results["resume.docx"]["email"] == 1

    redacted_document = docx.Document(str(output_dir / "resume.docx"))
    header_text = redacted_document.sections[0].header.paragraphs[0].text
    assert "jane.doe@example.com" not in header_text
    assert "[EMAIL]" in header_text


def test_round_trip_every_pii_class_leaves_zero_findings(tmp_path: Path):
    """The gate test: redact a fixture containing every PII class in both
    supported formats, then scan the redacted output and require zero
    findings — this is exactly the check B1.3b's redactor was missing."""
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    lines = [
        "Jane Doe",
        "jane.doe@example.com",
        "mailto:jane.doe2@example.com",
        "+91 98765 43210",
        "linkedin.com/in/janedoe",
        "linkedin/janedoe2",
        "github.com/janedoe",
        "github/janedoe2",
        "https://example.com/portfolio",
        "janedoe.dev",
    ]

    _make_pdf(input_dir / "resume.pdf", lines)
    _make_docx(input_dir / "resume.docx", lines)

    dry_run_results = redact_directory(input_dir, output_dir, dry_run=True)
    for file_name in ("resume.pdf", "resume.docx"):
        assert sum(dry_run_results[file_name].values()) == len(lines) - 1  # "Jane Doe" has no PII

    redact_directory(input_dir, output_dir, dry_run=False)

    rescan_results = scan_directory(output_dir)
    for file_name in ("resume.pdf", "resume.docx"):
        assert sum(rescan_results[file_name].values()) == 0, rescan_results[file_name]


def test_round_trip_names_and_pii_together_leaves_zero_findings(tmp_path: Path):
    """Same gate as above, but with --names-file-style name redaction on
    top: a fixture with a name (full form and first-name-alone) plus every
    PII class, in both formats, must scan clean once redacted with the
    compiled name pattern passed through."""
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    lines = [
        "Nick Miller",
        "Nick led the payments migration.",
        "Miller's team shipped on time.",
        "nick.miller@example.com",
        "+91 98765 43210",
        "linkedin.com/in/nickmiller",
        "github.com/nickmiller",
        "https://example.com/portfolio",
    ]

    _make_pdf(input_dir / "prac_001.pdf", lines)
    _make_docx(input_dir / "prac_001.docx", lines)

    name_patterns = compile_name_patterns({"prac_001": "Nick Miller"})

    redact_directory(input_dir, output_dir, dry_run=False, name_patterns=name_patterns)

    rescan_results = scan_directory(output_dir, name_patterns)
    for file_name in ("prac_001.pdf", "prac_001.docx"):
        assert sum(rescan_results[file_name].values()) == 0, rescan_results[file_name]


def test_name_redaction_is_scoped_to_its_own_resume(tmp_path: Path):
    """THE scoping test: a global name list would match common tokens
    ("Black") in every resume, not just the one they belong to. prac_001
    is named "Nick Black" and should have that name redacted; prac_002
    merely mentions "Black Friday pricing" and must be left alone."""
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    _make_docx(input_dir / "prac_001.docx", ["Nick Black", "Nick led the team."])
    _make_docx(
        input_dir / "prac_002.docx", ["Experienced retailer.", "Worked on Black Friday pricing."]
    )

    name_patterns = compile_name_patterns({"prac_001": "Nick Black"})

    redact_directory(input_dir, output_dir, dry_run=False, name_patterns=name_patterns)

    prac_001 = docx.Document(str(output_dir / "prac_001.docx"))
    prac_001_text = "\n".join(p.text for p in prac_001.paragraphs)
    assert "Nick Black" not in prac_001_text
    assert "Nick" not in prac_001_text
    assert "[NAME]" in prac_001_text

    prac_002 = docx.Document(str(output_dir / "prac_002.docx"))
    prac_002_text = "\n".join(p.text for p in prac_002.paragraphs)
    assert "Black Friday pricing" in prac_002_text
    assert "[NAME]" not in prac_002_text


def test_page_exceeding_rect_guard_stops_and_reports(tmp_path: Path):
    """A page that yields an implausible number of redaction rects means a
    pattern is over-matching; redacting it silently would destroy the
    document, so the redactor must stop instead of proceeding."""
    input_dir = tmp_path / "in"
    output_dir = tmp_path / "out"
    input_dir.mkdir()

    pdf_path = input_dir / "prac_001.pdf"
    repeat_count = MAX_REDACT_RECTS_PER_PAGE + 1
    doc = pymupdf.open()
    # A tall custom page so every inserted line stays inside the mediabox —
    # text extraction clips to the page rect, and a default-sized page is
    # nowhere near tall enough for 200+ lines.
    page = doc.new_page(width=595, height=72 + repeat_count * 10)
    for i in range(repeat_count):
        page.insert_text((72, 72 + i * 10), "Black")
    doc.save(pdf_path)
    doc.close()

    name_patterns = compile_name_patterns({"prac_001": "Black"})

    with pytest.raises(RedactionOvermatchError, match="prac_001"):
        redact_directory(input_dir, output_dir, dry_run=False, name_patterns=name_patterns)

"""Tests for scripts/strip_images.py.

All fixtures are built in-memory with PyMuPDF / zipfile — never read from
resumes/, resumes_clean/, or resumes_stripped/.
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import docx
import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

from strip_images import main, strip_directory, strip_docx_images, strip_pdf_images


def _red_png() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), False)
    pix.set_rect(pix.irect, (255, 0, 0))
    return pix.tobytes("png")


def _make_pdf(path: Path, *, with_image: bool, with_text: bool) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    if with_image:
        page.insert_image(pymupdf.Rect(10, 10, 60, 60), stream=_red_png())
    if with_text:
        page.insert_text((100, 100), "Experienced Python and SQL developer.")
    doc.save(path)
    doc.close()


_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Default Extension="png" ContentType="image/png"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

_ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

_DOCUMENT_XML_WITH_IMAGE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<w:body>
<w:p><w:r><w:t>Experienced Python and SQL developer.</w:t></w:r></w:p>
<w:p><w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">
<pic:blipFill><a:blip r:embed="rId2"/></pic:blipFill>
</pic:pic>
</a:graphicData>
</a:graphic>
</wp:inline></w:drawing></w:r></w:p>
</w:body>
</w:document>"""

_DOCUMENT_XML_NO_IMAGE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
<w:p><w:r><w:t>Experienced Python and SQL developer.</w:t></w:r></w:p>
</w:body>
</w:document>"""

_DOCUMENT_XML_IMAGE_ONLY = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<w:body>
<w:p><w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">
<pic:blipFill><a:blip r:embed="rId2"/></pic:blipFill>
</pic:pic>
</a:graphicData>
</a:graphic>
</wp:inline></w:drawing></w:r></w:p>
</w:body>
</w:document>"""

_DOCUMENT_RELS_WITH_IMAGE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>
</Relationships>"""


def _make_docx(path: Path, *, with_image: bool, with_text: bool = True) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("_rels/.rels", _ROOT_RELS)
        if with_image:
            document_xml = _DOCUMENT_XML_WITH_IMAGE if with_text else _DOCUMENT_XML_IMAGE_ONLY
            archive.writestr("word/document.xml", document_xml)
            archive.writestr("word/_rels/document.xml.rels", _DOCUMENT_RELS_WITH_IMAGE)
            archive.writestr("word/media/image1.png", _red_png())
        else:
            archive.writestr("word/document.xml", _DOCUMENT_XML_NO_IMAGE)


def test_strips_image_preserves_text(tmp_path: Path):
    in_path = tmp_path / "in.pdf"
    out_path = tmp_path / "out.pdf"
    _make_pdf(in_path, with_image=True, with_text=True)

    images_before, images_after, text_chars = strip_pdf_images(in_path, out_path)

    assert images_before == 1
    assert images_after == 0
    assert text_chars > 0

    with pymupdf.open(out_path) as doc:
        assert sum(len(page.get_images(full=True)) for page in doc) == 0
        assert "Python" in "".join(page.get_text() for page in doc)


def test_blank_after_strip_flag(tmp_path: Path, capsys: pytest.CaptureFixture[str]):
    in_dir = tmp_path / "in"
    out_dir = tmp_path / "out"
    in_dir.mkdir()
    _make_pdf(in_dir / "photo_only.pdf", with_image=True, with_text=False)

    exit_code = strip_directory(in_dir, out_dir)
    captured = capsys.readouterr()

    assert exit_code == 0
    assert "blank_after_strip: photo_only.pdf" in captured.out


def test_no_images_passes_through(tmp_path: Path, capsys: pytest.CaptureFixture[str]):
    in_dir = tmp_path / "in"
    out_dir = tmp_path / "out"
    in_dir.mkdir()
    _make_pdf(in_dir / "text_only.pdf", with_image=False, with_text=True)

    exit_code = strip_directory(in_dir, out_dir)
    captured = capsys.readouterr()

    assert exit_code == 0
    assert "blank_after_strip: text_only.pdf" not in captured.out
    assert "removed=0 remaining=0" in captured.out


def test_in_dir_equal_to_out_dir_refused(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    in_dir = tmp_path / "same"
    in_dir.mkdir()
    _make_pdf(in_dir / "resume.pdf", with_image=True, with_text=True)

    monkeypatch.setattr(
        sys, "argv", ["strip_images.py", "--in-dir", str(in_dir), "--out-dir", str(in_dir)]
    )

    exit_code = main()

    assert exit_code != 0
    with pymupdf.open(in_dir / "resume.pdf") as doc:
        assert sum(len(page.get_images(full=True)) for page in doc) == 1


def test_strips_docx_media_preserves_text(tmp_path: Path):
    in_path = tmp_path / "in.docx"
    out_path = tmp_path / "out.docx"
    _make_docx(in_path, with_image=True)

    media_before, media_after, text_chars = strip_docx_images(in_path, out_path)

    assert media_before == 1
    assert media_after == 0
    assert text_chars > 0

    with zipfile.ZipFile(out_path) as archive:
        assert not any(name.startswith("word/media/") for name in archive.namelist())


def test_stripped_docx_still_opens_via_docx_text_path(tmp_path: Path):
    in_path = tmp_path / "in.docx"
    out_path = tmp_path / "out.docx"
    _make_docx(in_path, with_image=True)

    strip_docx_images(in_path, out_path)

    document = docx.Document(str(out_path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Python" in text


def test_docx_with_no_media_passes_through_unchanged(tmp_path: Path):
    in_path = tmp_path / "in.docx"
    out_path = tmp_path / "out.docx"
    _make_docx(in_path, with_image=False)

    media_before, media_after, text_chars = strip_docx_images(in_path, out_path)

    assert media_before == 0
    assert media_after == 0
    assert text_chars > 0

    document = docx.Document(str(out_path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Python" in text


def test_docx_blank_after_strip_flag(tmp_path: Path, capsys: pytest.CaptureFixture[str]):
    in_dir = tmp_path / "in"
    out_dir = tmp_path / "out"
    in_dir.mkdir()
    _make_docx(in_dir / "photo_only.docx", with_image=True, with_text=False)

    exit_code = strip_directory(in_dir, out_dir)
    captured = capsys.readouterr()

    assert exit_code == 0
    assert "blank_after_strip: photo_only.docx" in captured.out
